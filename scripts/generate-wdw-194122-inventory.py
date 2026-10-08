#!/usr/bin/env python3
"""Generate WDW-194122 AOP pricing legend change inventory (HTML + PDF + XLSX)."""

from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "WDW-194122-pricing-legend-inventory"
ITEMS = Path.home() / "Documents/disney-dining-content-ops-full/data/dscribe/crawl/items.jsonl"

BASE = "https://dpep-dscribe-production.tridion.sdlproducts.com"
PUB = "914"
PUB_LABEL = "LGCY065 Parent (All) Publish"
CONTENT_PUB_LABEL = "LGCY045 Parent (All) Publish"

TICKET = "WDW-194122"
PARENT = "WDW-194121"
LAUNCH = "October 6, 2026"

ROWS_META = [
    (1, "Leaping Horse Libations", "https://disneyworld.disney.go.com/dining/boardwalk/leaping-horse-libations/", 1018516, "page", "general.constraints"),
    (2, "Satu'li Canteen", "https://disneyworld.disney.go.com/dining/animal-kingdom/satuli-canteen/", 1018754, "page", "general.metadata"),
    (3, "Cape Town Lounge and Wine Bar", "https://disneyworld.disney.go.com/dining/animal-kingdom-lodge/cape-town-lounge-and-wine-bar/", 1018454, "page", "general.constraints"),
    (4, "Sanaa Lounge", "https://disneyworld.disney.go.com/dining/animal-kingdom-villas-kidani/sanaa-lounge/", 1019514, "page", "general.constraints"),
    (5, "Uzima Springs Pool Bar", "https://disneyworld.disney.go.com/dining/animal-kingdom-lodge/uzima-springs-pool-bar/", 1018460, "page", "general.constraints"),
    (6, "The Hollywood Brown Derby Lounge", "https://disneyworld.disney.go.com/dining/hollywood-studios/hollywood-brown-derby-lounge/", 1019238, "page", "general.constraints"),
    (7, "Coral Reef Restaurant", "https://disneyworld.disney.go.com/dining/epcot/coral-reef-restaurant/", 1019590, "page", "general.constraints"),
    (8, "Le Cellier Steakhouse", "https://disneyworld.disney.go.com/dining/epcot/le-cellier-steakhouse/", 1019596, "page", "general.constraints"),
    (9, "Regal Eagle Smokehouse: Craft Drafts & Barbecue", "https://disneyworld.disney.go.com/dining/epcot/regal-eagle-smokehouse/", 1018910, "page", "general.constraints"),
    (10, "Crockett's Tavern", "https://disneyworld.disney.go.com/dining/cabins-at-fort-wilderness-resort/crockett-tavern/", 1018574, "page", "general.constraints"),
    (11, "Citricos Lounge", "https://disneyworld.disney.go.com/dining/grand-floridian-resort-and-spa/citricos-lounge/", 1018580, "page", "general.constraints"),
    (12, "Prince Eric's Village Market", "https://disneyworld.disney.go.com/dining/magic-kingdom/prince-eric-village-market/", 1019354, "page", "general.constraints"),
    (13, "Kona Island", "https://disneyworld.disney.go.com/dining/polynesian-resort/kona-island/", 1018630, "page", "general.constraints"),
    (14, "Bar Riva", "https://disneyworld.disney.go.com/dining/riviera-resort/bar-riva/", 1018651, "page", "general.constraints"),
    (15, "Crew's Cup Lounge", "https://disneyworld.disney.go.com/dining/yacht-club-resort/crew-cup-lounge/", 1018726, "page", "general.constraints"),
    (16, "EPCOT International Festival of the Arts – DISNEY ON BROADWAY Concert Series Dining Packages", "https://disneyworld.disney.go.com/dining/epcot/broadway-concert-series-dining-package/", 1552567, "comp", "general.content"),
    (17, "EPCOT International Festival of the Holidays – Candlelight Processional Dining Packages", "https://disneyworld.disney.go.com/dining/epcot/candlelight-dinner-packages/", 1328861, "comp", "general.content"),
    (19, "Holiday Kitchens at the Epcot International Festival of the Holidays", "https://disneyworld.disney.go.com/dining/epcot/holiday-kitchens/", 1019196, "page", "general.constraints"),
]


def load_crawl() -> dict[str, dict]:
    by_id: dict[str, dict] = {}
    if not ITEMS.is_file():
        return by_id
    with ITEMS.open(encoding="utf-8") as f:
        for line in f:
            try:
                by_id[json.loads(line)["id"]] = json.loads(line)
            except (json.JSONDecodeError, KeyError):
                continue
    return by_id


def to_pub(tcm: str) -> str:
    body = tcm.split(":", 1)[1]
    parts = body.split("-")
    if len(parts) >= 3:
        return f"tcm:{PUB}-{parts[1]}-{parts[2]}"
    return f"tcm:{PUB}-{parts[1]}"


def get_record(by_id: dict, num: int, kind: str) -> dict | None:
    if kind == "page":
        return by_id.get(f"tcm:627-{num}-64")
    return by_id.get(f"tcm:622-{num}")


def folder_chain(by_id: dict, parent_id: str | None) -> list[dict]:
    chain: list[dict] = []
    cur = parent_id
    seen: set[str] = set()
    while cur and cur not in seen:
        seen.add(cur)
        rec = by_id.get(cur)
        if not rec:
            break
        chain.append(rec)
        p = rec.get("parent")
        if not p or str(p).startswith("tcm:0-"):
            break
        cur = p
    chain.reverse()
    return chain


def explorer_url(by_id: dict, rec: dict, item_tcm_914: str) -> str:
    chain = folder_chain(by_id, rec.get("parent"))
    parts = [f"cme:publications_tcm:0-{PUB}-1"]
    if rec.get("type") == "Page":
        parts.append(f"tcm:{PUB}-3-4")
    for c in chain:
        parts.append(to_pub(c["id"]))
    container = "_".join(parts)
    return f"{BASE}/ui/explorer?container={container}&item={item_tcm_914}&panel=information"


def editor_url(num: int, kind: str, tab: str) -> str:
    if kind == "page":
        item = f"tcm:{PUB}-{num}-64"
        return f"{BASE}/ui/editor/page?activeItem={item}&item={item}&tab={tab}"
    item = f"tcm:{PUB}-{num}"
    return f"{BASE}/ui/editor/component?item={item}&tab={tab}"


def folder_location(rec: dict) -> str:
    path = rec.get("path", "")
    if rec.get("type") == "Page":
        return path.replace("LGCY060 US (en-US) Web Structure", PUB_LABEL)
    parent_path = path.rsplit("\\", 1)[0] if "\\" in path else path
    return parent_path.replace("LGCY040 US (en-US) Content", CONTENT_PUB_LABEL)


def build_rows(by_id: dict) -> list[dict]:
    rows: list[dict] = []
    for num, loc, prod, item_num, kind, tab in ROWS_META:
        rec = get_record(by_id, item_num, kind)
        item_914 = f"tcm:{PUB}-{item_num}-64" if kind == "page" else f"tcm:{PUB}-{item_num}"
        cms_name = rec.get("title", "—") if rec else "—"
        parent = by_id.get(rec["parent"]) if rec else None
        folder_short = parent.get("title", "—") if parent else "—"
        folder_path = folder_location(rec) if rec else "—"
        folder_link = explorer_url(by_id, rec, item_914) if rec else ""
        note = ""
        if num == 5 and cms_name != loc:
            note = "Verify: ticket links to page titled “Sanaa Lounge”; prod URL is Uzima Springs Pool Bar."
        if num == 4 and "Sanaa" in cms_name and "Lounge" in loc:
            note = "Verify: CMS page title is “Sanaa” (TSR); prod URL is Kidani Sanaa Lounge."
        rows.append(
            {
                "num": num,
                "location": loc,
                "prod_url": prod,
                "dscribe_url": editor_url(item_num, kind, tab),
                "folder_short": folder_short,
                "folder_path": folder_path,
                "folder_url": folder_link,
                "cms_name": cms_name,
                "item_uri": item_914,
                "note": note,
            }
        )
    return rows


def esc(s: str) -> str:
    return (
        s.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def write_html(rows: list[dict], path: Path) -> None:
    trs = []
    for r in rows:
        note_cell = f'<div class="note">{esc(r["note"])}</div>' if r["note"] else ""
        trs.append(
            f"""<tr>
  <td class="num">{r["num"]}</td>
  <td>{esc(r["location"])}{note_cell}</td>
  <td><a href="{esc(r["prod_url"])}">Prod</a></td>
  <td><a href="{esc(r["dscribe_url"])}">Edit in D-Scribe</a><br><span class="mono">{esc(r["item_uri"])}</span></td>
  <td><a href="{esc(r["folder_url"])}">{esc(r["folder_short"])}</a><br><span class="path">{esc(r["folder_path"])}</span></td>
  <td>{esc(r["cms_name"])}</td>
</tr>"""
        )
    body = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8"/>
  <title>{TICKET} — Pricing legend change inventory</title>
  <style>
    :root {{ font-family: system-ui, -apple-system, Segoe UI, Roboto, sans-serif; color: #1a1a1a; }}
    body {{ max-width: 1200px; margin: 2rem auto; padding: 0 1.5rem; line-height: 1.45; }}
    h1 {{ font-size: 1.35rem; margin-bottom: 0.25rem; }}
    .meta {{ color: #444; margin-bottom: 1.5rem; font-size: 0.95rem; }}
    .meta a {{ color: #0063e5; }}
    table {{ border-collapse: collapse; width: 100%; font-size: 0.82rem; }}
    th, td {{ border: 1px solid #ccc; padding: 0.45rem 0.5rem; vertical-align: top; }}
    th {{ background: #f0f4f8; text-align: left; }}
    tr:nth-child(even) td {{ background: #fafafa; }}
    a {{ color: #0063e5; word-break: break-all; }}
    .num {{ text-align: center; width: 2rem; }}
    .mono, .path {{ font-size: 0.72rem; color: #555; word-break: break-all; }}
    .note {{ color: #b45309; font-size: 0.75rem; margin-top: 0.35rem; }}
    @media print {{
      body {{ margin: 0.5in; }}
      a {{ color: #0063e5; text-decoration: underline; }}
    }}
  </style>
</head>
<body>
  <h1>FY27 AOP — Pricing legend update inventory</h1>
  <p class="meta">
    Jira: <a href="https://disneyexperiences.atlassian.net/browse/{TICKET}">{TICKET}</a>
    (parent <a href="https://disneyexperiences.atlassian.net/browse/{PARENT}">{PARENT}</a>) ·
    Target launch: <strong>{LAUNCH}</strong> · Generated {date.today().isoformat()} ·
    <strong>Do not publish early.</strong>
  </p>
  <table>
    <thead>
      <tr>
        <th>#</th>
        <th>Location</th>
        <th>Prod</th>
        <th>D-Scribe (edit)</th>
        <th>CMS folder</th>
        <th>Page / component</th>
      </tr>
    </thead>
    <tbody>
{"".join(trs)}
    </tbody>
  </table>
  <p class="meta" style="margin-top:1rem;">
    Tip: Open this file in a browser and use <strong>Print → Save as PDF</strong> to keep clickable links.
    For Teams, attach the <code>.html</code>, <code>.pdf</code>, or <code>.xlsx</code> from the same folder.
  </p>
</body>
</html>"""
    path.write_text(body, encoding="utf-8")


def write_xlsx(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    headers = ["#", "Location", "Prod URL", "D-Scribe edit URL", "CMS folder URL", "CMS folder path", "Page / component", "Notes"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(
            [
                r["num"],
                r["location"],
                r["prod_url"],
                r["dscribe_url"],
                r["folder_url"],
                r["folder_path"],
                r["cms_name"],
                r["note"],
            ]
        )
    for row_idx in range(2, ws.max_row + 1):
        for col in (3, 4, 5):
            cell = ws.cell(row=row_idx, column=col)
            url = cell.value
            if url and str(url).startswith("http"):
                cell.hyperlink = url
                cell.font = Font(color="0563C1", underline="single")
        for col in range(1, 9):
            ws.cell(row=row_idx, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    widths = [4, 36, 28, 36, 36, 48, 28, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)


def pdf_text(s: str) -> str:
    """Helvetica core font is Latin-1 only."""
    s = (
        s.replace("\u2014", "-")
        .replace("\u2013", "-")
        .replace("\u2019", "'")
        .replace("\u2018", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
        .replace("\u2026", "...")
    )
    return s.encode("latin-1", "replace").decode("latin-1")


def write_pdf(rows: list[dict], path: Path) -> None:
    from fpdf import FPDF

    class PDF(FPDF):
        def header(self):
            self.set_font("Helvetica", "B", 11)
            self.cell(0, 8, pdf_text(f"{TICKET} - Pricing legend change inventory (launch {LAUNCH})"), ln=True)
            self.set_font("Helvetica", "", 8)
            self.set_text_color(80, 80, 80)
            self.cell(0, 5, "Do not publish early. Links require Disney network/VPN.", ln=True)
            self.set_text_color(0, 0, 0)
            self.ln(2)

        def footer(self):
            self.set_y(-12)
            self.set_font("Helvetica", "I", 8)
            self.cell(0, 8, f"Page {self.page_no()}/{{nb}}", align="C")

    pdf = PDF(orientation="L", unit="mm", format="A4")
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()
    pdf.set_font("Helvetica", "", 7)

    col_w = [8, 52, 18, 55, 55, 45, 38]
    headers = ["#", "Location", "Prod", "D-Scribe", "CMS folder", "CMS item", "Notes"]
    pdf.set_fill_color(240, 244, 248)
    pdf.set_font("Helvetica", "B", 7)
    for h, w in zip(headers, col_w):
        pdf.cell(w, 6, h, border=1, fill=True)
    pdf.ln()

    pdf.set_font("Helvetica", "", 6.5)
    line_h = 4

    def multi_link_cell(w, h, lines: list[tuple[str, str | None]], border=1):
        x, y = pdf.get_x(), pdf.get_y()
        pdf.rect(x, y, w, h)
        pdf.set_xy(x + 1, y + 1)
        for i, (text, url) in enumerate(lines):
            if i:
                pdf.ln(line_h - 1)
            if url:
                pdf.set_text_color(0, 99, 229)
                pdf.write(line_h, text, link=url)
                pdf.set_text_color(0, 0, 0)
            else:
                pdf.multi_cell(w - 2, line_h, text)
        pdf.set_xy(x + w, y)

    for r in rows:
        x0 = pdf.get_x()
        y0 = pdf.get_y()
        row_h = max(
            14,
            6 + (len(r["location"]) // 42) * 4,
            6 + (len(r["note"]) // 38) * 4 if r["note"] else 0,
        )
        if y0 + row_h > 190:
            pdf.add_page()
            y0 = pdf.get_y()

        pdf.set_xy(x0, y0)
        pdf.multi_cell(col_w[0], row_h, str(r["num"]), border=1)
        pdf.set_xy(x0 + col_w[0], y0)
        loc_text = pdf_text(r["location"] + ("\n" + r["note"] if r["note"] else ""))
        pdf.multi_cell(col_w[1], line_h, loc_text, border=1)
        pdf.set_xy(x0 + col_w[0] + col_w[1], y0)
        pdf.cell(col_w[2], row_h, "Open", border=1, link=r["prod_url"])
        pdf.set_xy(x0 + col_w[0] + col_w[1] + col_w[2], y0)
        pdf.cell(col_w[3], row_h, "Edit", border=1, link=r["dscribe_url"])
        pdf.set_xy(x0 + col_w[0] + col_w[1] + col_w[2] + col_w[3], y0)
        pdf.cell(col_w[4], row_h, pdf_text(r["folder_short"])[:28], border=1, link=r["folder_url"])
        pdf.set_xy(x0 + col_w[0] + col_w[1] + col_w[2] + col_w[3] + col_w[4], y0)
        pdf.multi_cell(col_w[5], line_h, pdf_text(r["cms_name"]), border=1)
        pdf.set_xy(x0 + sum(col_w[:6]), y0)
        pdf.multi_cell(col_w[6], line_h, pdf_text(r["note"] or "-"), border=1)
        pdf.set_xy(x0, y0 + row_h)

    pdf.output(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    by_id = load_crawl()
    rows = build_rows(by_id)
    write_html(rows, OUT / f"{TICKET}-inventory.html")
    write_xlsx(rows, OUT / f"{TICKET}-inventory.xlsx")
    write_pdf(rows, OUT / f"{TICKET}-inventory.pdf")
    print(f"Wrote {OUT}/")


if __name__ == "__main__":
    main()
